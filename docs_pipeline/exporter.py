"""Export Markdown / structured content to DOCX and PDF with optional Verification Summary."""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
    HRFlowable,
)


def _set_cell_shading(cell, color_hex: str) -> None:
    """Set background color on a table cell (docx)."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)


def _status_color(status: str) -> str:
    s = (status or "").upper()
    if "VERIFIED" in s or "SUPPORTED" in s or "PASS" in s:
        return "C6EFCE"
    if "PARTIAL" in s or "REVIEW" in s or "WARN" in s:
        return "FFEB9C"
    if "FAIL" in s or "UNSUPPORTED" in s or "FAKE" in s or "NO_" in s:
        return "FFC7CE"
    return "D9E2F3"


def _build_verification_section_docx(doc: Document, verification: Dict[str, Any]) -> None:
    """Append a Verification Summary section to a DOCX document."""
    doc.add_page_break()
    _add_heading(doc, "Verification Summary", level=1)

    meta = doc.add_paragraph()
    meta.add_run(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}").italic = True
    meta.add_run(f"\nDocument ID: {verification.get('document_id', '—')}")
    meta.add_run(f"\nOverall Status: {verification.get('overall_status', 'UNKNOWN')}")

    if verification.get("notice"):
        notice = doc.add_paragraph()
        run = notice.add_run(verification["notice"])
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    citations: List[Dict] = verification.get("citations") or []
    if not citations:
        doc.add_paragraph("No citations were extracted or verified.")
        return

    # Summary table
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["#", "Citation", "Match Status", "Score", "Source"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
        _set_cell_shading(cell, "1A56DB")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)

    for idx, c in enumerate(citations, 1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = (c.get("original_text") or c.get("text") or "")[:120]
        status = c.get("match_status") or c.get("status") or c.get("label") or "—"
        row[2].text = str(status)
        score = c.get("score") or c.get("confidence") or "—"
        row[3].text = f"{score:.2f}" if isinstance(score, float) else str(score)
        src = c.get("matched_url") or c.get("source_url") or "—"
        row[4].text = str(src)[:60]
        _set_cell_shading(row[2], _status_color(str(status)))

    doc.add_paragraph()
    # Sources used
    sources = verification.get("sources") or []
    if sources:
        _add_heading(doc, "Sources Consulted", level=2)
        for s in sources:
            p = doc.add_paragraph(style="List Bullet")
            title = s.get("title") or s.get("url") or "Unknown"
            p.add_run(f"{title} — {s.get('url', '')} [{s.get('fetch_status', '')}]")


def export_docx(
    content: str,
    output: Union[str, Path, None] = None,
    title: str = "DraftMate Document",
    verification: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Export Markdown-ish content to DOCX.
    Simple parser: headings (#), paragraphs, lists.
    Optionally append Verification Summary.
    """
    doc = Document()
    # Title
    t = doc.add_heading(title, 0)
    for run in t.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    # Simple markdown → docx
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", stripped):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", stripped), style="List Number")
        else:
            # paragraph; handle **bold** lightly
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        i += 1

    if verification:
        _build_verification_section_docx(doc, verification)

    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if output:
        Path(output).write_bytes(data)
    return data


def _build_verification_story(verification: Dict[str, Any], styles) -> list:
    """Build ReportLab flowables for verification summary."""
    story = []
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1A56DB")))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Verification Summary", styles["Heading1Custom"]))
    story.append(
        Paragraph(
            f"<i>Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}</i><br/>"
            f"Document ID: {verification.get('document_id', '—')}<br/>"
            f"Overall Status: <b>{verification.get('overall_status', 'UNKNOWN')}</b>",
            styles["Body"],
        )
    )
    if verification.get("notice"):
        story.append(Paragraph(f"<font size='8' color='grey'>{verification['notice']}</font>", styles["Body"]))

    citations = verification.get("citations") or []
    if not citations:
        story.append(Paragraph("No citations were extracted or verified.", styles["Body"]))
        return story

    # Table data
    data = [["#", "Citation (excerpt)", "Status", "Score", "Source"]]
    for idx, c in enumerate(citations, 1):
        status = str(c.get("match_status") or c.get("status") or c.get("label") or "—")
        score = c.get("score") or c.get("confidence") or "—"
        score_s = f"{score:.2f}" if isinstance(score, float) else str(score)
        excerpt = (c.get("original_text") or c.get("text") or "")[:80]
        src = str(c.get("matched_url") or c.get("source_url") or "—")[:40]
        data.append([str(idx), excerpt, status, score_s, src])

    col_widths = [0.4 * inch, 3.2 * inch, 1.3 * inch, 0.7 * inch, 1.8 * inch]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A56DB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 0), (3, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for r, c in enumerate(citations, 1):
        status = str(c.get("match_status") or c.get("status") or c.get("label") or "")
        hex_c = _status_color(status)
        style_cmds.append(("BACKGROUND", (2, r), (2, r), colors.HexColor(f"#{hex_c}")))
    t.setStyle(TableStyle(style_cmds))
    story.append(Spacer(1, 8))
    story.append(t)

    sources = verification.get("sources") or []
    if sources:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Sources Consulted", styles["Heading2Custom"]))
        for s in sources:
            title = s.get("title") or s.get("url") or "Unknown"
            story.append(
                Paragraph(
                    f"• {title} — <font size='7'>{s.get('url', '')}</font> [{s.get('fetch_status', '')}]",
                    styles["Body"],
                )
            )
    return story


def export_pdf(
    content: str,
    output: Union[str, Path, None] = None,
    title: str = "DraftMate Document",
    verification: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Export content to PDF with optional Verification Summary page."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="Heading1Custom",
            parent=styles["Heading1"],
            textColor=colors.HexColor("#1A56DB"),
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Heading2Custom",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#1A56DB"),
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            spaceAfter=6,
            alignment=TA_JUSTIFY,
        )
    )

    story = []
    story.append(Paragraph(title, styles["Heading1Custom"]))
    story.append(Spacer(1, 6))

    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue
        if stripped.startswith("# "):
            story.append(Paragraph(stripped[2:], styles["Heading1Custom"]))
        elif stripped.startswith("## "):
            story.append(Paragraph(stripped[3:], styles["Heading2Custom"]))
        elif stripped.startswith("### "):
            story.append(Paragraph(stripped[4:], styles["Heading2"]))
        else:
            # escape basic HTML entities
            text = (
                stripped.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace("**", "")
            )
            story.append(Paragraph(text, styles["Body"]))

    if verification:
        story.append(PageBreak())
        story.extend(_build_verification_story(verification, styles))

    doc.build(story)
    data = buf.getvalue()
    if output:
        Path(output).write_bytes(data)
    return data


def export_with_verification(
    content: str,
    verification: Dict[str, Any],
    formats: List[str] = ("docx", "pdf"),
    title: str = "DraftMate Document",
    output_dir: Union[str, Path] = ".",
) -> Dict[str, bytes]:
    """
    Export content + verification summary to requested formats.
    Returns dict of format → bytes.
    """
    out: Dict[str, bytes] = {}
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    if "docx" in formats:
        data = export_docx(content, verification=verification, title=title)
        out["docx"] = data
        Path(output_dir, f"{title.replace(' ', '_')}.docx").write_bytes(data)
    if "pdf" in formats:
        data = export_pdf(content, verification=verification, title=title)
        out["pdf"] = data
        Path(output_dir, f"{title.replace(' ', '_')}.pdf").write_bytes(data)
    return out
